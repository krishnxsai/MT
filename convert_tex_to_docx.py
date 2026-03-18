#!/usr/bin/env python3
"""
Extract content from meditrack_paper.tex and create a DOCX document
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import re

def tex_to_docx(tex_file, docx_file):
    """Convert TEX to DOCX"""

    # Read the TEX file
    with open(tex_file, 'r', encoding='utf-8') as f:
        content = f.read()

    # Create DOCX document
    doc = Document()

    # Extract and add title
    title_match = re.search(r'\\title\{([^}]+)\}', content)
    if title_match:
        title_text = title_match.group(1).replace('\n', ' ').replace('\\', '')
        title = doc.add_heading(title_text, level=0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Extract author info
    author_section = re.search(r'\\author\{(.*?)\n\n', content, re.DOTALL)
    if author_section:
        author_lines = author_section.group(1)
        # Clean up and add authors
        author_text = re.sub(r'\\thanks\{', '', author_lines)
        author_text = re.sub(r'\}\s*\n?\s*\\thanks', '\n', author_text)
        author_text = author_text.rstrip('}')
        author_text = re.sub(r'\$\^{\d+}\$', '', author_text)  # Remove superscripts

        for line in author_text.split('\n'):
            if line.strip() and not line.strip().startswith('%'):
                p = doc.add_paragraph(line.strip())
                p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    # Extract abstract
    abstract_match = re.search(r'\\begin\{abstract\}(.*?)\\end\{abstract\}', content, re.DOTALL)
    if abstract_match:
        doc.add_heading('Abstract', level=1)
        abstract_text = abstract_match.group(1).strip()
        # Clean up LaTeX commands
        abstract_text = re.sub(r'\\textit\{([^}]*)\}', r'\1', abstract_text)
        abstract_text = re.sub(r'\\textbf\{([^}]*)\}', r'\1', abstract_text)
        abstract_text = re.sub(r'\\tt\s*', '', abstract_text)
        abstract_text = re.sub(r'~', ' ', abstract_text)
        abstract_text = re.sub(r'\s+', ' ', abstract_text).strip()

        doc.add_paragraph(abstract_text)

    # Extract sections (I, II, III, etc)
    sections = re.findall(r'\\section\{([^}]*)\}(.*?)(?=\\section|\\end\{document\})', content, re.DOTALL)

    for section_title, section_content in sections:
        if section_title.strip():
            doc.add_heading(section_title.strip(), level=1)

        # Extract paragraphs
        paragraphs = re.split(r'\n\s*\n', section_content)
        for para in paragraphs:
            para = para.strip()
            if para and not para.startswith('%'):
                # Clean LaTeX commands
                para = re.sub(r'\\textit\{([^}]*)\}', r'\1', para)
                para = re.sub(r'\\textbf\{([^}]*)\}', r'\1', para)
                para = re.sub(r'\\tt\s*', '', para)
                para = re.sub(r'\\cite\{[^}]*\}', '[citation]', para)
                para = re.sub(r'~', ' ', para)
                para = re.sub(r'\s+', ' ', para).strip()

                if para and not para.startswith('\\'):
                    doc.add_paragraph(para)

    # Save document
    doc.save(docx_file)
    print(f"✓ Successfully created: {docx_file}")

if __name__ == '__main__':
    tex_file = r"c:/Users/wujjw/Music/MediTrack-main/MediTrack-main/alternate_tj_latex_template_ap/meditrack_paper.tex"
    docx_file = r"c:/Users/wujjw/Music/MediTrack-main/MediTrack-main/alternate_tj_latex_template_ap/meditrack_paper.docx"

    tex_to_docx(tex_file, docx_file)
